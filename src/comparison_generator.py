"""
Comparison Generator Module

Builds comparison tables from matched results and exports to various formats.
"""

from typing import List, Dict
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from src.semantic_matcher import MatchResult


class ComparisonGenerator:
    """Generates comparison tables from match results."""

    def __init__(self):
        """Initialize ComparisonGenerator."""
        pass

    def build_comparison_table(self,
                              base_data: pd.DataFrame,
                              comparison_data: Dict[str, pd.DataFrame],
                              match_results: Dict[str, List[MatchResult]]) -> pd.DataFrame:
        """
        Build comparison table from matched results.

        Args:
            base_data: Normalized base DataFrame with [item, unit, price]
            comparison_data: Dict mapping filename to normalized DataFrame
            match_results: Dict mapping filename to list of MatchResult

        Returns:
            DataFrame with comparison table
        """
        # Start with base data
        comparison_rows = []

        for base_idx, base_row in base_data.iterrows():
            row_data = {
                'Base_Item': base_row['item'],
                'Base_Unit': base_row['unit'],
                'Base_Price': base_row['price']
            }

            # Add comparison data for each file
            for file_name, comp_df in comparison_data.items():
                # Find match result for this base item
                matches = match_results.get(file_name, [])

                # Find match for current base index
                match = next((m for m in matches if m.base_index == base_idx), None)

                # Create column prefix
                col_prefix = self._get_column_prefix(file_name)

                if match and match.match_index >= 0:
                    # Get matched row from comparison data
                    comp_row = comp_df.iloc[match.match_index]

                    row_data[f'{col_prefix}_Item'] = comp_row['item']
                    row_data[f'{col_prefix}_Unit'] = comp_row['unit']
                    row_data[f'{col_prefix}_Price'] = comp_row['price']
                    row_data[f'{col_prefix}_Match'] = match.confidence
                    row_data[f'{col_prefix}_Similarity'] = match.similarity_score
                else:
                    # No match found
                    row_data[f'{col_prefix}_Item'] = 'NO ENCONTRADO'
                    row_data[f'{col_prefix}_Unit'] = '-'
                    row_data[f'{col_prefix}_Price'] = None
                    row_data[f'{col_prefix}_Match'] = 'NONE'
                    row_data[f'{col_prefix}_Similarity'] = 0.0

            comparison_rows.append(row_data)

        return pd.DataFrame(comparison_rows)

    def _get_column_prefix(self, filename: str) -> str:
        """Generate column prefix from filename."""
        # Remove extension and clean up
        name = Path(filename).stem
        # Limit length
        if len(name) > 15:
            name = name[:15]
        # Remove special characters
        name = ''.join(c if c.isalnum() else '_' for c in name)
        return name

    def add_statistics_row(self, df: pd.DataFrame) -> pd.DataFrame:
        """
        Add summary statistics row at the bottom.

        Args:
            df: Comparison DataFrame

        Returns:
            DataFrame with statistics row added
        """
        stats_row = {'Base_Item': 'ESTADÍSTICAS'}

        # Calculate match rates for each comparison file
        for col in df.columns:
            if col.endswith('_Match'):
                # Count matches (not NONE)
                total = len(df)
                matched = len(df[df[col] != 'NONE'])
                match_rate = (matched / total * 100) if total > 0 else 0
                stats_row[col] = f"{matched}/{total} ({match_rate:.1f}%)"
            elif col not in ['Base_Item']:
                stats_row[col] = '-'

        stats_df = pd.DataFrame([stats_row])
        return pd.concat([df, stats_df], ignore_index=True)

    def export_csv(self, df: pd.DataFrame, filepath: Path) -> bool:
        """
        Export comparison table to CSV.

        Args:
            df: Comparison DataFrame
            filepath: Output file path

        Returns:
            True if successful, False otherwise
        """
        try:
            # Ensure exports directory exists
            filepath.parent.mkdir(parents=True, exist_ok=True)

            # Export to CSV
            df.to_csv(filepath, index=False, encoding='utf-8-sig')

            print(f"✓ Exported to CSV: {filepath}")
            return True

        except Exception as e:
            print(f"Error exporting to CSV: {e}")
            return False

    def export_excel(self, df: pd.DataFrame, filepath: Path,
                    add_formatting: bool = True) -> bool:
        """
        Export comparison table to Excel with optional formatting.

        Args:
            df: Comparison DataFrame
            filepath: Output file path
            add_formatting: Whether to add colors and formatting

        Returns:
            True if successful, False otherwise
        """
        try:
            # Ensure exports directory exists
            filepath.parent.mkdir(parents=True, exist_ok=True)

            # Export to Excel
            df.to_excel(filepath, index=False, engine='openpyxl')

            if add_formatting:
                self._add_excel_formatting(filepath)

            print(f"✓ Exported to Excel: {filepath}")
            return True

        except Exception as e:
            print(f"Error exporting to Excel: {e}")
            return False

    def _add_excel_formatting(self, filepath: Path):
        """Add formatting to Excel file."""
        try:
            wb = load_workbook(filepath)
            ws = wb.active

            # Define colors
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)

            high_match_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            medium_match_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
            low_match_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            no_match_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

            # Format header row
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            # Find match columns
            match_columns = []
            for col_idx, cell in enumerate(ws[1], 1):
                if cell.value and str(cell.value).endswith('_Match'):
                    match_columns.append(col_idx)

            # Apply conditional formatting based on match confidence
            for row_idx in range(2, ws.max_row + 1):
                for col_idx in match_columns:
                    cell = ws.cell(row=row_idx, column=col_idx)
                    value = cell.value

                    if value == 'HIGH':
                        cell.fill = high_match_fill
                    elif value == 'MEDIUM':
                        cell.fill = medium_match_fill
                    elif value == 'LOW':
                        cell.fill = low_match_fill
                    elif value == 'NONE':
                        cell.fill = no_match_fill

            # Auto-adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter

                for cell in column:
                    try:
                        if cell.value:
                            max_length = max(max_length, len(str(cell.value)))
                    except:
                        pass

                adjusted_width = min(max_length + 2, 50)  # Cap at 50
                ws.column_dimensions[column_letter].width = adjusted_width

            # Save formatted workbook
            wb.save(filepath)

        except Exception as e:
            print(f"Warning: Could not add Excel formatting: {e}")

    def export_markdown(self, df: pd.DataFrame, filepath: Path) -> bool:
        """
        Export comparison table to Markdown format.

        Args:
            df: Comparison DataFrame
            filepath: Output file path

        Returns:
            True if successful, False otherwise
        """
        try:
            # Ensure exports directory exists
            filepath.parent.mkdir(parents=True, exist_ok=True)

            # Convert to markdown
            markdown_table = df.to_markdown(index=False)

            # Write to file
            with open(filepath, 'w', encoding='utf-8') as f:
                f.write("# Tabla de Comparación de Presupuestos\n\n")
                f.write(markdown_table)
                f.write("\n\n---\n")
                f.write(f"\n*Generado con Gemini File Search Chat*\n")

            print(f"✓ Exported to Markdown: {filepath}")
            return True

        except Exception as e:
            print(f"Error exporting to Markdown: {e}")
            return False

    def generate_summary_report(self, base_data: pd.DataFrame,
                               comparison_data: Dict[str, pd.DataFrame],
                               match_results: Dict[str, List[MatchResult]]) -> str:
        """
        Generate text summary report of comparison.

        Args:
            base_data: Base DataFrame
            comparison_data: Dict of comparison DataFrames
            match_results: Dict of match results

        Returns:
            Formatted summary string
        """
        report = "\n" + "="*60 + "\n"
        report += "RESUMEN DE COMPARACIÓN\n"
        report += "="*60 + "\n\n"

        report += f"Archivo base: {len(base_data)} ítems\n\n"

        for file_name, comp_df in comparison_data.items():
            report += f"Archivo: {file_name}\n"
            report += f"  Ítems totales: {len(comp_df)}\n"

            # Get matches for this file
            matches = match_results.get(file_name, [])

            # Calculate statistics
            total = len(matches)
            high = sum(1 for m in matches if m.confidence == 'HIGH')
            medium = sum(1 for m in matches if m.confidence == 'MEDIUM')
            low = sum(1 for m in matches if m.confidence == 'LOW')
            none = sum(1 for m in matches if m.confidence == 'NONE')

            matched = high + medium + low
            match_rate = (matched / total * 100) if total > 0 else 0

            report += f"  Coincidencias:\n"
            report += f"    - Alta confianza: {high}\n"
            report += f"    - Media confianza: {medium}\n"
            report += f"    - Baja confianza: {low}\n"
            report += f"    - No encontrado: {none}\n"
            report += f"  Tasa de coincidencia: {match_rate:.1f}%\n\n"

        report += "="*60 + "\n"

        return report
